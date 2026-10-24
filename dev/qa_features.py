"""
dev/qa_features.py — feature-by-feature QA of the dashboard, run against the
STAND-IN executable (not DLBreach). It checks behaviour, not numbers.

For each of the four worked examples (T1–T4) it drives the app's own
buttons, uploads and downloads, in the local version and in the browser
version, and checks what comes back. It also checks the project and
template files, autosave and restore, stale-result flags, and the security
settings. Every check is reported; the script carries on after a failure
and exits non-zero if any check failed.

Linux and macOS only (the stand-in cannot run on Windows).

    python dev/qa_features.py
"""

from __future__ import annotations

import io
import json
import os
import re
import subprocess
import sys
import tempfile
import tomllib
import traceback
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "dev"))
os.chdir(ROOT)

import pandas as pd                    # noqa: E402

import test_app_headless as H          # noqa: E402  installs the Streamlit stand-in
from make_standin import make_exe      # noqa: E402
from dlbreach_io import (DEFAULT_PARAMS, build_input_text, parse_out,  # noqa: E402
                         read_input_file, run_case)

ST = H.ST
SS = ST.session_state
APP = (ROOT / "app.py").read_text(encoding="utf-8")
CASES = ["T1_water_overtopping", "T2_water_piping",
         "T3_water_overtopping", "T4_water_piping"]
PAGES = ["1. About", "2. Project", "3. Model setup", "4. Single run",
         "5. Sensitivity scenarios", "6. Monte Carlo", "7. Case histories",
         "8. Export"]
RESULTS: list[tuple[str, str, bool, str]] = []
EXE = None


class Upload:
    """What st.file_uploader returns."""

    def __init__(self, name: str, data: bytes):
        self.name, self._d, self.size = name, data, len(data)

    def getvalue(self) -> bytes:
        return self._d


UPLOADS: dict[str, object] = {}
DOWNLOADS: dict[str, tuple[bytes, str]] = {}


def render(press=(), browser=False) -> dict:
    """One script run. Buttons in `press` (label or key) return True;
    uploads come from UPLOADS; downloads are captured in DOWNLOADS."""
    ob, of, od = H._Stub.button, H._Stub.file_uploader, H._Stub.download_button

    def button(self, label="", *a, key=None, **k):
        self._claim(key)
        if k.get("disabled"):
            return False                               # as in Streamlit
        return label in press or (key is not None and key in press)

    def file_uploader(self, label="", *a, key=None, **k):
        self._claim(key)
        return UPLOADS.get(key)

    def download_button(self, label="", data=None, *a, file_name=None, **k):
        if data is None:
            data = k.get("data")
        if isinstance(data, str):
            data = data.encode()
        DOWNLOADS[label] = (data, file_name)
        return False

    H._Stub.button, H._Stub.file_uploader, H._Stub.download_button = \
        button, file_uploader, download_button
    plat = sys.platform
    sys.platform = "emscripten" if browser else "linux"
    guards = _browser_guards() if browser else None
    try:
        ST.calls.clear()
        ST.html_out.clear()
        ST.specs.clear()
        ST.md_out.clear()
        ST.keys_seen.clear()
        g = {"__name__": "__main__", "__file__": str(ROOT / "app.py")}
        exec(compile(APP, "app.py", "exec"), g)
        H._raise_contained()
        return g
    finally:
        sys.platform = plat
        if guards:
            guards()
        H._Stub.button, H._Stub.file_uploader, H._Stub.download_button = ob, of, od


BROWSER_VIOLATIONS: list[str] = []


def _browser_guards():
    """In browser mode the app must not start processes or write files in
    the dashboard folder (Pyodide has no real processes; anything written
    would not persist and must not be relied on). Record any attempt."""
    import builtins
    import io as _io
    real_open, real_popen = builtins.open, subprocess.Popen

    def guarded_open(file, mode="r", *a, **k):
        if any(c in mode for c in "wax+"):
            try:
                path = Path(file).resolve()
                if ROOT in path.parents or path == ROOT:
                    BROWSER_VIOLATIONS.append(f"write {path.relative_to(ROOT)}")
            except TypeError:
                pass
        return real_open(file, mode, *a, **k)

    class GuardedPopen(real_popen):
        def __init__(self, args, *a, **k):
            BROWSER_VIOLATIONS.append(f"process {args!r}")
            super().__init__(args, *a, **k)

    builtins.open = _io.open = guarded_open
    subprocess.Popen = GuardedPopen

    def restore():
        builtins.open = _io.open = real_open
        subprocess.Popen = real_popen
    return restore


class Skipped(Exception):
    """A check that could not run here; reported as SKIP, not as a pass."""


def check(area: str, name: str):
    """Decorator: run a check, record pass, fail or skip, never stop."""
    def wrap(fn):
        try:
            detail = fn() or ""
            RESULTS.append((area, name, True, str(detail)))
        except Skipped as exc:
            RESULTS.append((area, name, None, f"skipped: {exc}"))
        except Exception as exc:                        # noqa: BLE001
            tb = traceback.format_exc(limit=3).strip().splitlines()[-1]
            RESULTS.append((area, name, False, f"{type(exc).__name__}: {exc} | {tb}"))
        return fn
    return wrap


def fresh(page="1. About"):
    SS.clear()
    UPLOADS.clear()
    DOWNLOADS.clear()
    SS["page"] = page
    SS["exe_path"] = str(EXE)
    SS["_session_started"] = True
    SS["_restore_done"] = True


def load_example(case: str, browser=False):
    SS["page"] = PAGES[1]
    SS["keep_example"] = case
    render(browser=browser)
    render(press=("load_example",), browser=browser)
    render(browser=browser)


def standin_outputs(inputs: dict[str, str]) -> bytes:
    """Run the stand-in on (case -> input text); return a zip of .out files,
    as a user would after run_all.bat."""
    buf = io.BytesIO()
    with tempfile.TemporaryDirectory() as td, zipfile.ZipFile(buf, "w") as z:
        for name, text in inputs.items():
            r = run_case(EXE, text, case_name=name, work_dir=td)
            assert r["ok"], f"stand-in failed on {name}"
            z.write(r["out_path"], f"{name}.out")
    return buf.getvalue()


def zip_inputs(data: bytes) -> dict[str, str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        assert "run_all.bat" in names and "README.txt" in names, names
        return {Path(n).stem: z.read(n).decode().replace("\r\n", "\n")
                for n in names if n.endswith(".txt") and n != "README.txt"}


# ---------------------------------------------------------------------------
def case_checks(case: str):
    data = json.loads((ROOT / "test_cases" / f"{case}.json").read_text())
    area = case[:2]

    @check(area, "input file round trip (write, read, write)")
    def _():
        p = dict(DEFAULT_PARAMS, **data)
        p["res_curve"] = [tuple(x) for x in p["res_curve"]]
        t1 = build_input_text(p)
        p2 = dict(DEFAULT_PARAMS, **read_input_file(t1))
        p2["case_name"], p2["notes"] = p["case_name"], p["notes"]
        assert build_input_text(p2) == t1
        txt = (ROOT / "test_cases" / f"{case}.txt").read_text()
        assert build_input_text(p) == txt, "stored .txt differs from the .json"

    fresh()
    load_example(case)

    @check(area, "load worked example")
    def _():
        assert SS["params"]["case_name"] == data["case_name"]
        assert SS["params"]["breach_mode"] == data["breach_mode"]

    @check(area, "Model setup in model, site and geodetic datums")
    def _():
        for d in ("model", "site", "geo"):
            SS["params"]["datum_display"] = d
            SS["page"] = PAGES[2]
            g = render()
            assert not g["water_checks"](), g["water_checks"]()
        SS["params"]["datum_display"] = data.get("datum_display", "model")

    @check(area, "Model setup: input file download matches the writer")
    def _():
        SS["page"] = PAGES[2]
        render()
        b, fn = DOWNLOADS["Download input file"]
        assert b.decode() == build_input_text(SS["params"])

    # --- local single run
    @check(area, "Single run (local): run, results, downloads")
    def _():
        SS["page"] = PAGES[3]
        render(press=("Run DLBreach",))
        df = SS["result_df"]
        assert df is not None and len(df) > 10 and df["Q_breach"].max() > 0
        g = render()
        assert g["result_state"]() == "current"
        h, fn = DOWNLOADS["Download hydrograph for inundation modelling (.csv)"]
        hy = pd.read_csv(io.BytesIO(h), comment="#")
        assert len(hy) > 10, hy.shape
        geo, _ = DOWNLOADS["Download breach geometry (.txt)"]
        assert len(geo) > 100
        return f"peak (stand-in) {df['Q_breach'].max():,.0f} m3/s, {len(hy)} csv rows"

    @check(area, "stale-result flag after an input change")
    def _():
        SS["params"]["sed_d50"] = SS["params"]["sed_d50"] * 2
        g = render()
        assert g["result_state"]() == "stale"
        SS["params"]["sed_d50"] = SS["params"]["sed_d50"] / 2
        assert render()["result_state"]() == "current"

    @check(area, "Sensitivity scenarios (local): run all, results CSV")
    def _():
        SS["page"] = PAGES[4]
        render()
        n = int(SS.get("keep_scen_n", 3))
        render(press=(f"Run all {n} runs",))
        sc = SS["scen_store"]
        assert sc and len(sc["summary"]) == n
        render()
        r, _ = DOWNLOADS["Download results (.csv)"]
        assert len(pd.read_csv(io.BytesIO(r))) == n
        return f"{n} runs"

    @check(area, "Monte Carlo (local): 20 realisations, downloads")
    def _():
        SS["page"] = PAGES[5]
        SS["mc_nreal"] = 20
        render()
        render(press=("Run Monte Carlo",))
        res = next(iter(SS["mc_store"].values()))["results"]
        assert int(res["ok"].sum()) == 20, res[["ok", "error"]].head()
        render()
        for lab in ("Download summary (.csv)", "Input file (.txt)",
                    "Hydrograph (.csv)", "All realisations (.csv)"):
            assert lab in DOWNLOADS and len(DOWNLOADS[lab][0]) > 0, lab

    @check(area, "Case histories with the run")
    def _():
        SS["page"] = PAGES[6]
        calls = render()
        assert ST.calls.count("vega_lite_chart") >= 1

    @check(area, "Export: PDF and Word open and name the case")
    def _():
        SS["page"] = PAGES[7]
        for k in ("mc", "mc_env", "mc_drivers", "cases", "scenarios"):
            SS[f"keep_rsec_{k}"] = True
        out = []
        for fmt in ("PDF", "Word"):
            SS["keep_rep_fmt"] = fmt
            render()
            render(press=(f"Build {fmt}",))
            data_, ext, _ = SS["report_file"]
            if fmt == "PDF":
                assert data_[:5] == b"%PDF-"
                pages = len(re.findall(rb"/Type\s*/Page[^s]", data_))
                assert pages >= 3, pages
                out.append(f"PDF {pages} pages")
            else:
                import docx
                d = docx.Document(io.BytesIO(data_))
                text = "\n".join(p.text for p in d.paragraphs)
                assert SS["params"]["case_name"] in text or len(d.tables) > 0
                out.append(f"Word {len(d.paragraphs)} paragraphs, {len(d.tables)} tables")
        return "; ".join(out)

    # --- browser version: download inputs, run outside, upload outputs
    fresh()
    load_example(case, browser=True)

    @check(area, "browser: no executable needed, Project shows the browser note")
    def _():
        g = render(browser=True)
        assert g["BROWSER"] and g["exe_status"]() == ("", False)

    @check(area, "browser Single run: download input, upload .out")
    def _():
        SS["page"] = PAGES[3]
        render(browser=True)
        b, fn = DOWNLOADS["1 · Download input file"]
        name = Path(fn).stem
        z = standin_outputs({name: b.decode()})
        with zipfile.ZipFile(io.BytesIO(z)) as zz:
            UPLOADS["out_loader"] = Upload(f"{name}.out", zz.read(f"{name}.out"))
        render(browser=True)
        assert SS["result_df"] is not None and len(SS["result_df"]) > 10
        assert render(browser=True)["result_state"]() == "loaded"
        UPLOADS.pop("out_loader")

    @check(area, "browser Sensitivity scenarios: zip out, outputs back")
    def _():
        SS["page"] = PAGES[4]
        DOWNLOADS.clear()
        render(browser=True)
        lab = next(k for k in DOWNLOADS if k.startswith("1 · Download") and "input files" in k)
        inputs = zip_inputs(DOWNLOADS[lab][0])
        UPLOADS["scen_out_up"] = [Upload("outputs.zip", standin_outputs(inputs))]
        render(browser=True)
        sc = SS["scen_store"]
        assert sc and len(sc["summary"]) == len(inputs)
        assert len(sc["hydros"]) == len(inputs), "not every run matched its output"
        UPLOADS.pop("scen_out_up")
        return f"{len(inputs)} runs matched by file name"

    @check(area, "browser Monte Carlo: zip out, outputs back")
    def _():
        SS["page"] = PAGES[5]
        SS["mc_nreal"] = 20
        DOWNLOADS.clear()
        render(browser=True)
        lab = next(k for k in DOWNLOADS if k.startswith("1 · Download") and "input files" in k)
        inputs = zip_inputs(DOWNLOADS[lab][0])
        assert len(inputs) == 20, len(inputs)
        UPLOADS["mc_out_up"] = [Upload("mc.zip", standin_outputs(inputs))]
        render(browser=True)
        res = next(iter(SS["mc_store"].values()))["results"]
        assert int(res["ok"].sum()) == 20, int(res["ok"].sum())
        UPLOADS.pop("mc_out_up")

    @check(area, "browser Export: PDF and Word")
    def _():
        SS["page"] = PAGES[7]
        for fmt in ("PDF", "Word"):
            SS["keep_rep_fmt"] = fmt
            render(browser=True)
            render(press=(f"Build {fmt}",), browser=True)
            assert SS["report_file"][0][:2] in (b"%P", b"PK")


# ---------------------------------------------------------------------------
def project_checks():
    area = "Project"
    fresh()
    load_example("T2_water_piping")

    @check(area, "Save project, then open it again")
    def _():
        SS["page"] = PAGES[1]
        render()
        b, fn = DOWNLOADS["Save project (.json)"]
        saved = json.loads(b)
        SS["params"]["emb_height"] = 99.0
        UPLOADS["proj_loader"] = Upload(fn, b)
        render()
        UPLOADS.pop("proj_loader")
        render()
        assert SS["params"]["emb_height"] == saved["emb_height"] == 12.0
        assert SS["params"]["res_curve"][-1] == tuple(saved["res_curve"][-1])

    @check(area, "Load an existing DLBreach input file (.txt)")
    def _():
        txt = (ROOT / "test_cases" / "T4_water_piping.txt").read_bytes()
        UPLOADS["txt_loader"] = Upload("T4_water_piping.txt", txt)
        render()
        UPLOADS.pop("txt_loader")
        assert SS["params"]["breach_mode"] == 2 and SS["params"]["pipe_depth"] == 9.0
        assert SS["params"]["sediment_type"] == 2

    @check(area, "Templates zip: five files, input template reads back")
    def _():
        render()
        b, _ = DOWNLOADS["Download all templates (.zip)"]
        with zipfile.ZipFile(io.BytesIO(b)) as z:
            names = set(z.namelist())
            assert {"README.txt", "dlbreach_input_template.txt",
                    "stage_storage_template.csv", "scenario_runs_template.csv",
                    "project_template.json"} <= names, names
            p = read_input_file(z.read("dlbreach_input_template.txt").decode())
            assert "emb_height" in p
            pd.read_csv(io.BytesIO(z.read("stage_storage_template.csv")), comment="#")

    @check(area, "Model setup: upload a stage-storage CSV and use it")
    def _():
        SS["page"] = PAGES[2]
        SS["params"]["ss_source"] = "file"
        SS["ss_source"] = "file"
        csv = b"Elevation (m),Volume (m3)\n0,0\n4,500000\n8,2000000\n12,5000000\n14,7000000\n"
        UPLOADS["ss_upload"] = Upload("ss.csv", csv)
        g = render()
        SS["params"]["datum_base_rl"] = 0.0
        render(press=("Use this stage-storage",))
        UPLOADS.pop("ss_upload")
        rc = SS["params"]["res_curve"]
        assert len(rc) == 5 and rc[-1][1] == 7e6, rc

    @check(area, "Scenarios: the downloaded table (CSV and Excel) uploads back")
    def _():
        import dlbreach_batch as B
        for fmt in ("CSV", "Excel"):
            SS["page"] = PAGES[4]
            SS.pop("scen_sig", None)
            render()
            b, fn = DOWNLOADS[f"Download as {fmt}"]
            # fill in three D50 values (mm), as a user would
            tab = pd.read_csv(io.BytesIO(b)) if fmt == "CSV" else pd.read_excel(io.BytesIO(b))
            runs = [c for c in tab.columns if c not in ("Parameter", "Unit")]
            row = tab["Parameter"] == "Median grain size D50"
            for r, v in zip(runs, (0.5, 2.0, 8.0)):
                tab[r] = tab[r].astype(object)
                tab.loc[row, r] = v
            out = io.BytesIO()
            (tab.to_csv(out, index=False) if fmt == "CSV" else tab.to_excel(out, index=False))
            UPLOADS["scen_upload"] = Upload(fn, out.getvalue())
            render()
            UPLOADS.pop("scen_upload")
            assert SS["keep_scen_keys"] == ["sed_d50"], SS.get("keep_scen_keys")
            vals, rep = B.read_grid(SS["scen_grid"])
            got = [round(vals[r]["sed_d50"] * 1000, 6) for r in rep["runs"]]
            assert got == [0.5, 2.0, 8.0], got
            assert not rep["unrecognised"], rep
        return "values 0.5, 2, 8 mm read back from both"

    @check(area, "Scenarios: the template from the templates zip uploads")
    def _():
        import dlbreach_batch as B
        SS["page"] = PAGES[1]
        render()
        b, _ = DOWNLOADS["Download all templates (.zip)"]
        with zipfile.ZipFile(io.BytesIO(b)) as z:
            t = z.read("scenario_runs_template.csv")
        g2 = B.read_grid_file(t, "scenario_runs_template.csv")
        vals, rep = B.read_grid(g2)
        assert rep["runs"] and not rep["unrecognised"], rep
        return f"{len(rep['runs'])} runs, {sum(len(v) for v in vals.values())} values"

    @check(area, "Autosave, then restore in a new session (local only)")
    def _():
        auto = ROOT / ".dlbreach_autosave.json"
        auto.unlink(missing_ok=True)
        SS["page"] = PAGES[1]
        SS["params"]["case_name"] = "autosave_check"
        render()
        assert auto.is_file(), "autosave not written"
        # a new session: restore offer, then press Restore
        SS.clear()
        SS["page"] = PAGES[0]
        render()
        render(press=("restore_yes",))
        render()
        assert SS["params"]["case_name"] == "autosave_check"
        auto.unlink()

    @check(area, "browser: nothing written to disk (no autosave)")
    def _():
        auto = ROOT / ".dlbreach_autosave.json"
        auto.unlink(missing_ok=True)
        fresh(PAGES[1])
        SS["params"] = dict(DEFAULT_PARAMS, case_name="browser_check")
        render(browser=True)
        render(browser=True)
        assert not auto.exists()


def published_files(env) -> list[Path]:
    """The files git would publish from this folder, per .gitignore."""
    with tempfile.TemporaryDirectory() as td:
        idx = dict(env, GIT_INDEX_FILE=str(Path(td) / "index"))
        out = subprocess.run(["git", "--git-dir", str(Path(td) / "g"), "--work-tree",
                              str(ROOT), "init", "-q"], env=idx, check=True)
        r = subprocess.run(["git", "--git-dir", str(Path(td) / "g"), "--work-tree",
                            str(ROOT), "ls-files", "--others", "--exclude-standard"],
                           env=idx, check=True, capture_output=True, text=True)
    return [ROOT / x for x in r.stdout.splitlines() if x]


# ---------------------------------------------------------------------------
def security_checks():
    area = "Security"

    @check(area, "local server: localhost only, no usage statistics")
    def _():
        cfg = tomllib.loads((ROOT / ".streamlit" / "config.toml").read_text())
        assert cfg["server"]["address"] == "localhost"
        assert cfg["browser"]["gatherUsageStats"] is False

    @check(area, ".gitignore blocks executables, outputs, project data")
    def _():
        with tempfile.TemporaryDirectory() as td:
            env = dict(os.environ, GIT_CONFIG_COUNT="1",
                       GIT_CONFIG_KEY_0="safe.directory", GIT_CONFIG_VALUE_0="*")
            subprocess.run(["git", "init", "-q", td], check=True, env=env)
            (Path(td) / ".gitignore").write_text((ROOT / ".gitignore").read_text())
            for d in ("test_cases", "dev/fixtures", ".streamlit"):
                (Path(td) / d).mkdir(parents=True, exist_ok=True)
            must = ["DLBreach.exe", "Run4_v3.txt", "Run4_v3.out", "case.out",
                    "SampleRunLog.xlsx", "data.xlsx", "x.xlsm", ".dlbreach_autosave.json",
                    ".dlbreach_config.json", ".venv/x", "dlb_abc/x.txt",
                    "client_case.txt", "my_project.json", "hydrograph.csv", "runs.zip",
                    "report.docx", "report.pdf", ".streamlit/secrets.toml",
                    "test_private_benchmarks.py", "dev/private_terms.txt",
                    "test_cases/client.json", "dev/fixtures/other.out", "uploaded.out",
                    "test_cases/wu_2016/client.txt", "test_cases/wu_2016/Teton.out",
                    "test_cases/wu_2016/Teton_cards.json"]
            keep = ["dev/fixtures/standin_T1.out", "test_cases/T1_water_overtopping.txt",
                    "test_cases/T4_water_piping.json", "requirements.txt", "app.py",
                    "dlbreach_io.py", ".streamlit/config.toml", "dev/qa_features.py",
                    "docs/images/model-setup.png", "README.md",
                    "dev/scan_confidential_names.py", "dev/standin_dlbreach.py",
                    "docs/QA_REPORT.md",
                    "test_cases/wu_2016/Teton_cards.txt",
                    "test_cases/wu_2016/Sutter_cards_BasinVolume.txt",
                    "test_cases/wu_2016/Mecox_WaveDS.txt"]
            for f in must:
                r = subprocess.run(["git", "-C", td, "check-ignore", "-q", f], env=env)
                assert r.returncode == 0, f"not ignored: {f}"
            for f in keep:
                r = subprocess.run(["git", "-C", td, "check-ignore", "-q", f], env=env)
                assert r.returncode != 0, f"wrongly ignored: {f}"

    @check(area, "no confidential names in the published files")
    def _():
        from scan_confidential_names import scan
        res = scan()
        if res is None:
            raise Skipped("dev/private_terms.txt not present")
        assert not res[1], res[1][:5]
        return f"{res[0]} names, none found"

    @check(area, "browser mode: no processes started, nothing written to the folder")
    def _():
        assert not BROWSER_VIOLATIONS, BROWSER_VIOLATIONS[:5]
        return "checked during every browser-mode page run above"


def loading_checks():
    """Loading DLBreach input files: name and notes, card-for-card fidelity,
    Prof. Wu's validation inputs (when present locally) and data files."""
    from dlbreach_io import card_differences
    area = "Loading"

    @check(area, "input file sets the case name and notes")
    def _():
        fresh(PAGES[1])
        txt = (ROOT / "test_cases" / "T2_water_piping.txt").read_bytes()
        UPLOADS["txt_loader"] = Upload("T2_water_piping.txt", txt)
        render()
        UPLOADS.pop("txt_loader")
        render()
        assert SS["params"]["case_name"] == "T2_water_piping", SS["params"]["case_name"]
        assert SS["params"]["notes"].startswith("Test 2"), SS["params"]["notes"]
        r = SS["load_report"]
        assert not (r["missing"] or r["changed"]), r
        assert SS.get("result_df") is None             # results start afresh

    @check(area, "power-law reservoir: viewing the table changes nothing")
    def _():
        fresh(PAGES[2])
        SS["params"] = dict(DEFAULT_PARAMS, reservoir_method=3, res_VH=6.0e8,
                            res_H=24.5, res_m=2.0, case_name="powerlaw")
        render()
        before = DOWNLOADS["Download input file"][0]
        SS["keep_in_ss_source"] = "table"
        render()
        render()
        SS["keep_in_ss_source"] = "power"
        render()
        after = DOWNLOADS["Download input file"][0]
        assert before == after, "input file changed"
        assert b"Upstream_Reservoir 3," in after

    @check(area, "table reservoir: viewing the power-law inputs changes nothing")
    def _():
        fresh()
        load_example("T1_water_overtopping")
        SS["page"] = PAGES[2]
        render()
        before = DOWNLOADS["Download input file"][0]
        SS["keep_in_ss_source"] = "power"
        render()
        render()
        mid = DOWNLOADS["Download input file"][0]
        SS["keep_in_ss_source"] = "table"
        render()
        after = DOWNLOADS["Download input file"][0]
        assert before == mid == after, "input file changed"
        assert b"Upstream_Reservoir                   0, 27" in after

    @check(area, "loaded case switched to cohesive writes kd and tauc")
    def _():
        txt = (ROOT / "test_cases" / "T1_water_overtopping.txt").read_text()
        txt = "\n".join(ln for ln in txt.splitlines()
                        if not ln.startswith(("Cohesive_Soil_Erosion_kd",
                                              "Cohesive_Soil_Erosion_Tauc")))
        fresh(PAGES[1])
        UPLOADS["txt_loader"] = Upload("nokd.txt", txt.encode())
        render()
        UPLOADS.pop("txt_loader")
        SS["page"] = PAGES[2]
        render()
        assert b"Cohesive_Soil_Erosion_kd" not in DOWNLOADS["Download input file"][0]
        SS["params"]["sediment_type"] = 2
        SS["params"]["kd"] = 55.0
        SS.pop("keep_in_sediment_type", None)
        SS.pop("keep_in_kd", None)
        SS.pop("_model_keep_in_kd", None)
        render()
        w = DOWNLOADS["Download input file"][0]
        assert b"Cohesive_Soil_Erosion_kd 55.0" in w, w[-600:]

    wu_dir = ROOT / "test_cases" / "wu_2016"
    files = sorted(wu_dir.glob("*_cards*.txt")) if wu_dir.is_dir() else []
    if not files:
        @check(area, "Prof. Wu's validation inputs")
        def _():
            raise Skipped("test_cases/wu_2016 not present")
        return
    for f in files:
        @check(area, f"Wu case {f.stem}: load, write back, run")
        def _(f=f):
            fresh(PAGES[1])
            SS["keep_example"] = f"wu:{f.stem}"
            render()
            render(press=("load_example",))
            render()
            assert SS["params"]["case_name"] == f.stem
            SS["page"] = PAGES[2]
            render()
            written = DOWNLOADS["Download input file"][0].decode()
            d = card_differences(f.read_text(errors="replace"), written)
            assert not (d["missing"] or d["changed"] or d["added"]), d
            g = render()
            assert not g["water_checks"](), g["water_checks"]()
            SS["page"] = PAGES[3]
            render(press=("Run DLBreach",))
            assert SS["result_df"] is not None, "run gave no output"
            kept = len(SS["params"].get("extra_cards") or [])
            refs = SS["params"].get("data_file_refs") or []
            return f"{d['n']} cards identical, {kept} kept" + \
                (f", {len(refs)} data files attached" if refs else "")

    mecox = wu_dir / "Mecox_cards_WSLbay.txt"
    if mecox.is_file():
        @check(area, "data files: uploaded case blocks runs until they are added")
        def _():
            fresh(PAGES[1])
            UPLOADS["txt_loader"] = Upload(mecox.name, mecox.read_bytes())
            render()
            UPLOADS.pop("txt_loader")
            g = render()
            refs = SS["params"]["data_file_refs"]
            assert len(g["water_checks"]()) == len(refs) == 4, g["water_checks"]()
            # scenario and Monte Carlo runs are blocked too
            SS["page"] = PAGES[4]
            render()
            n = int(SS.get("keep_scen_n", 3))
            render(press=(f"Run all {n} runs",))
            assert not SS.get("scen_store"), "scenarios ran without data files"
            SS["page"] = PAGES[1]
            UPLOADS["data_files_up"] = [Upload(r, (wu_dir / r).read_bytes()) for r in refs]
            render()
            UPLOADS.pop("data_files_up")
            g = render()
            assert not g["water_checks"](), g["water_checks"]()
            # placed next to the input for a run, and in run packages
            with tempfile.TemporaryDirectory() as td:
                run_case(EXE, "x", case_name="c", work_dir=td,
                         data_files=SS["params"]["data_files"])
                assert all((Path(td) / r).is_file() for r in refs)
            SS["page"] = PAGES[4]
            DOWNLOADS.clear()
            render(browser=True)
            lab = next(k for k in DOWNLOADS if "input files" in k)
            with zipfile.ZipFile(io.BytesIO(DOWNLOADS[lab][0])) as z:
                assert set(refs) <= set(z.namelist()), z.namelist()
            return "4 files: runs blocked, attached, copied to runs and zips"


# ---------------------------------------------------------------------------
def _xs(press=()) -> str:
    """Render, and return the cross-section drawing shown."""
    render(press=press)
    return next((h for h in ST.html_out if "Dam Cross-Section" in h), "")


def _md() -> str:
    return "\n".join(ST.md_out)


def _load_wu(stem: str):
    fresh(PAGES[1])
    SS["keep_example"] = f"wu:{stem}"
    render()
    render(press=("load_example",))
    render()


def schematic_checks():
    import re
    from dlbreach_viz import C_CORE, C_COVER
    area = "Clay core and drawings"
    wu = ROOT / "test_cases" / "wu_2016"

    @check(area, "distance axis labelled on overtopping and piping sections")
    def _():
        out = []
        for case in ("T1_water_overtopping", "T2_water_piping"):
            fresh()
            load_example(case)
            SS["page"] = PAGES[2]
            svg = _xs()
            ticks = re.findall(r'text-anchor="middle" font-size="10.5" '
                               r'fill="#5E6B77">(-?[\d.,]+)</text>', svg)
            assert len(ticks) >= 3, (case, ticks)
            assert "Distance from upstream toe" in svg
            out.append(f"{case[:2]}: {len(ticks)} labels")
        return ", ".join(out)

    @check(area, "loaded clay core: read into inputs, drawn, not listed as kept")
    def _():
        _load_wu("Banqiao_cards")
        P = SS["params"]
        assert P["has_core"] and P["core_height"] == 23.0 and P["core_center"] == -3.0
        SS["page"] = PAGES[2]
        svg = _xs()
        assert C_CORE in svg and "Clay core height" in svg and "3 m upstream" in svg
        kept = [c.split(None, 1)[0] for c in P.get("extra_cards") or []]
        assert not any(k.startswith("Core_") for k in kept), kept
        assert "Upstream_Reservoir_Inflow" in kept
        assert SS["keep_in_has_core"] is True
        assert "Submode 3 needs" not in _md()
        return f"core 23 m, offset 3 m upstream; {len(kept)} card kept"

    @check(area, "core added to a case: written, drawn, removed again")
    def _():
        fresh()
        load_example("T1_water_overtopping")
        SS["page"] = PAGES[2]
        render()
        before = DOWNLOADS["Download input file"][0]
        assert b"Core_Height" not in before
        SS["keep_in_has_core"] = True
        render()
        svg = _xs()
        P = SS["params"]
        assert P["has_core"] and P["core_height"] == round(0.9 * P["emb_height"], 2)
        w = DOWNLOADS["Download input file"][0]
        assert f"Core_Height {P['core_height']!r}".encode() in w, w[-900:]
        assert C_CORE in svg
        assert "Used with overtopping submode 3" in _md()   # T1 is submode 1
        SS["keep_in_has_core"] = False
        render()
        render()
        assert DOWNLOADS["Download input file"][0] == before
        return "15 core cards written, then none"

    if (wu / "Impactfld03_cards.txt").is_file():
        @check(area, "loaded core: edited cards written, kd and tauc on going cohesive")
        def _():
            _load_wu("Impactfld03_cards")          # non-cohesive core, no kd or tauc
            SS["page"] = PAGES[2]
            render()
            w0 = DOWNLOADS["Download input file"][0]
            assert b"Core_Cohesive_Erosion_kd" not in w0
            SS["keep_in_core_sed_type"] = 2
            render()
            render()
            w1 = DOWNLOADS["Download input file"][0].decode()
            assert "Core_Noncohesive_or_Cohesive 2" in w1
            assert "Core_Cohesive_Erosion_kd" in w1 and "Core_Cohesive_Erosion_Tauc" in w1
            return "kd and tauc written once the core is cohesive"

    @check(area, "submode 3 without a core, and a core outside the dam, flagged")
    def _():
        fresh()
        load_example("T1_water_overtopping")
        SS["page"] = PAGES[2]
        SS["keep_in_overtopping_mode"] = 3
        render()
        render()
        assert "Submode 3 needs a clay core" in _md()
        assert "no core is described" in _md()
        SS["keep_in_has_core"] = True
        render()
        assert "Submode 3 needs a clay core" not in _md()   # same render
        SS["params"]["core_center"] = -40.0
        SS.pop("keep_in_core_center", None)
        SS.pop("_model_keep_in_core_center", None)
        render()
        assert "upstream face lies outside" in _md(), _md()[-500:]

    @check(area, "core inputs offered for scenarios only in submode 3")
    def _():
        _load_wu("Banqiao_cards")
        g = render()
        assert "core_kd" in g["scenario_keys_available"]()
        SS["params"]["overtopping_mode"] = 1
        assert not any(k.startswith("core_") for k in g["scenario_keys_available"]())

    if (wu / "Yahekou_cards.txt").is_file():
        @check(area, "clay cover from a loaded file drawn on the cross-section")
        def _():
            _load_wu("Yahekou_cards")
            SS["page"] = PAGES[2]
            svg = _xs()
            assert C_COVER in svg and "downstream slope, 0.3 m" in svg
            assert "shows the clay cover" in _md()

    if (wu / "Teton_cards.txt").is_file():
        @check(area, "one-sided breach drawn against a labelled abutment")
        def _():
            _load_wu("Teton_cards")
            SS["page"] = PAGES[3]
            render(press=("Run DLBreach",))
            render()
            face = next((h for h in ST.html_out if "Post-Breach Face" in h), "")
            assert "abutment (not erodible)" in face, "no abutment label"
            assert "One-sided (at abutment)" in face
            assert "widens to the right only" in face

        @check(area, "erosion limit below the storage table: a caution, not an error")
        def _():
            _load_wu("Teton_cards")
            SS["page"] = PAGES[2]
            render()
            md = _md()
            assert "Erosion limit below the stage-storage table" in md
            assert "lies below the base of the stage-storage" not in md


# ---------------------------------------------------------------------------
CASE_KEYS = ["c_qhb", "c_qvw", "c_bvw", "c_tvw"]


def _case_specs() -> list[dict]:
    """The four case-history chart specifications of the last render."""
    sp = [x for x in ST.specs
          if x["layer"] and "scatter" not in str(x.get("title", ""))
          and any("x" in v and "y" in v for v in x["data"]["values"][:1])]
    titles = [x["layer"][-1]["encoding"]["x"]["title"] for x in sp]
    assert len(sp) >= 4, (len(sp), titles)
    sp = sp[-4:]
    assert sorted(_case_key(x) for x in sp) == sorted(CASE_KEYS), titles
    return sp


def _case_key(spec: dict) -> str:
    """Which of the four charts a specification is, from its axis titles."""
    enc = spec["layer"][-1]["encoding"]
    x, y = enc["x"]["title"], enc["y"]["title"]
    if x.startswith("Breach height"):
        return "c_qhb"
    return {"Peak outflow (m³/s)": "c_qvw", "Average breach width (m)": "c_bvw",
            "Failure time (h)": "c_tvw"}[y]


def _series_counts(spec: dict) -> dict:
    out = {}
    for v in spec["data"]["values"]:
        out[v["Series"]] = out.get(v["Series"], 0) + 1
    return out


def overlay_checks(case: str):
    """Scenario and Monte Carlo runs on the four case-history charts."""
    import numpy as np
    from dlbreach_empirical import derive_inputs
    from dlbreach_mc import realisation_params
    area = f"Case-history overlays ({case[:2]})"
    fresh()
    load_example(case)
    SS["page"] = PAGES[6]
    render()

    @check(area, "before any runs: toggles off, nothing to add yet")
    def _():
        md = " ".join(ST.md_out)
        assert "Nothing to add yet from: Single run, Sensitivity scenarios, Monte Carlo" in md, md[-400:]

    SS["page"] = PAGES[3]
    render(press=("Run DLBreach",))
    SS["page"] = PAGES[4]
    render()
    n_scen = int(SS.get("keep_scen_n", 3))
    render(press=(f"Run all {n_scen} runs",))
    SS["page"] = PAGES[5]
    SS["mc_nreal"] = 20
    render()
    # vary the erosion limit and grain size so the realisations spread out
    SS["mc_chosen"] = ["hard_bottom_elev", "sed_d50"]
    render()
    SS["mc_d_hard_bottom_elev"] = SS["mc_d_sed_d50"] = "uniform"
    render()
    hb = float(SS["params"]["emb_height"])
    SS["mc_lo_hard_bottom_elev"], SS["mc_hi_hard_bottom_elev"] = 0.0, 0.6 * hb
    render()
    render(press=("Run Monte Carlo",))

    @check(area, "each scenario and realisation stores water volume and formation time")
    def _():
        case_ = SS["scen_store"]["case"]
        assert len(case_) == n_scen, len(case_)
        assert all(np.isfinite(v["water_volume_above_invert_m3"]) for v in case_.values())
        res = next(iter(SS["mc_store"].values()))["results"]
        ok = res[res["ok"] == True]                         # noqa: E712
        for c in ("water_volume_above_invert_m3", "formation_time_hr"):
            assert c in res.columns, c
        assert ok["water_volume_above_invert_m3"].notna().all()
        return (f"{n_scen} scenarios, {len(ok)} realisations; formation time "
                f"found for {int(ok['formation_time_hr'].notna().sum())}")

    @check(area, "stored values equal a fresh run of the same inputs")
    def _():
        ent = next(iter(SS["mc_store"].values()))
        res, cols, base = ent["results"], ent["cols"], ent["base"]
        checked = 0
        with tempfile.TemporaryDirectory() as td:
            for i in list(res.index[res["ok"] == True])[:3]:        # noqa: E712
                p = realisation_params(base, res, i, cols)
                r = run_case(EXE, build_input_text(p), case_name="chk", work_dir=td)
                d = derive_inputs(p, parse_out(r["out_path"]))
                for k, col in (("Vw", "water_volume_above_invert_m3"),
                               ("tf_hr", "formation_time_hr")):
                    a, b = d[k], res.loc[i, col]
                    if a is None:
                        assert pd.isna(b), (k, a, b)
                    else:
                        assert abs(a - b) <= 1e-9 * max(1.0, abs(a)), (k, a, b)
                checked += 1
            sc = SS["scen_store"]
            name = next(iter(sc["case"]))
            p = sc["params"][name]
            r = run_case(EXE, build_input_text(p), case_name="chk", work_dir=td)
            d = derive_inputs(p, parse_out(r["out_path"]))
            assert abs(d["Vw"] - sc["case"][name]["water_volume_above_invert_m3"]) \
                <= 1e-9 * max(1.0, d["Vw"])
        return f"{checked} realisations and scenario '{name}' recomputed"

    SS["page"] = PAGES[6]

    @check(area, "default: the single run only")
    def _():
        render()
        for sp in _case_specs():
            c = _series_counts(sp)
            assert c.get("This analysis") == 1, c
            assert not any(k.startswith(("Monte Carlo", "Scenario")) for k in c), c

    @check(area, "scenarios and Monte Carlo added to all four charts")
    def _():
        SS["keep_cases_scen"] = True
        SS["keep_cases_mc"] = True
        render()
        res = next(iter(SS["mc_store"].values()))["results"]
        ok = res[res["ok"] == True]                         # noqa: E712
        tw = ok["final_top_width_m"] + ok["final_bottom_width_m"]
        want = {"c_qhb": ((ok["max_breach_depth_m"] > 0) & (ok["peak_outflow_m3s"] > 0)),
                "c_qvw": ((ok["water_volume_above_invert_m3"] > 0) & (ok["peak_outflow_m3s"] > 0)),
                "c_bvw": ((ok["water_volume_above_invert_m3"] > 0) & (tw > 0)),
                "c_tvw": ((ok["water_volume_above_invert_m3"] > 0) & (ok["formation_time_hr"] > 0))}
        got = []
        for sp in _case_specs():
            key = _case_key(sp)
            c = _series_counts(sp)
            assert c.get("This analysis") == 1, c
            assert c.get("Monte Carlo", 0) == int(want[key].sum()), (key, c, int(want[key].sum()))
            assert 1 <= c.get("Scenario runs", 0) <= n_scen, (key, c)
            got.append(f"{c.get('Monte Carlo', 0)}+{c.get('Scenario runs', 0)}")
        hbs = ok["max_breach_depth_m"]
        assert hbs.max() - hbs.min() > 0.05 * hbs.max(), "ensemble did not spread"
        return "MC+scenario points per chart: " + ", ".join(got)

    @check(area, "realisations drawn small and light; scenarios and single run on top")
    def _():
        sp = _case_specs()[0]
        marks = [(l["mark"].get("size"), l["mark"].get("opacity"),
                  json.dumps(l.get("transform", ""))) for l in sp["layer"]]
        faint = [m for m in marks if m[0] == 22]
        assert faint and '"oneOf": ["Monte Carlo"]' in faint[0][2], marks
        assert marks.index(faint[0]) == 0, "realisations not drawn first"
        assert marks[-1][0] == 330, "single run not drawn last"
        leg = sp["layer"][-1]["encoding"]["color"]["scale"]["domain"]
        assert leg[-1] == "This analysis", leg

    @check(area, "hover names the run and realisation")
    def _():
        ev = {v["Event"] for v in _case_specs()[0]["data"]["values"]
              if v["Series"] in ("Monte Carlo", "Scenario runs")}
        assert any(", realisation " in e for e in ev), list(ev)[:3]
        assert set(SS["scen_store"]["case"]) & ev, list(ev)[:5]

    @check(area, "single run switched off removes the diamond")
    def _():
        SS["keep_cases_det"] = False
        render()
        for sp in _case_specs():
            assert "This analysis" not in _series_counts(sp)
        SS["keep_cases_det"] = True

    @check(area, "two ensembles: one series and colour each")
    def _():
        store = dict(SS["mc_store"])
        k0 = next(iter(store))
        SS["mc_store"] = {k0: store[k0], "Second": store[k0]}
        render()
        c = _series_counts(_case_specs()[0])
        SS["mc_store"] = store
        assert f"Monte Carlo: {k0}" in c and "Monte Carlo: Second" in c, c
        return ", ".join(f"{k} {v}" for k, v in c.items() if k.startswith("Monte"))

    @check(area, "results from before 1.0.3: left out and named")
    def _():
        store = SS["mc_store"]
        k0 = next(iter(store))
        old = dict(store[k0], results=store[k0]["results"].drop(
            columns=["water_volume_above_invert_m3", "formation_time_hr"]))
        SS["mc_store"] = {k0: old}
        render()
        md = " ".join(ST.md_out)
        c = _series_counts(_case_specs()[0])
        SS["mc_store"] = store
        assert "Some runs are not shown." in md and k0 in md
        assert not any(k.startswith("Monte Carlo") for k in c), c

    @check(area, "PDF and Word appendix carry the added runs")
    def _():
        SS["page"] = PAGES[7]
        SS["keep_rsec_cases"] = True
        out = []
        for fmt in ("PDF", "Word"):
            SS["keep_rep_fmt"] = fmt
            render()
            render(press=(f"Build {fmt}",))
            data_, _, _ = SS["report_file"]
            if fmt == "PDF":
                assert data_[:5] == b"%PDF-"
                out.append(f"PDF {len(data_) // 1024} kB")
            else:
                import docx
                d = docx.Document(io.BytesIO(data_))
                text = " ".join(p.text for p in d.paragraphs)
                assert "scenario runs are green triangles" in text, text[:300]
                assert "Monte Carlo realisations are small light points" in text
                out.append("Word note names both")
        SS["page"] = PAGES[6]
        return "; ".join(out)

    @check(area, "browser Monte Carlo (outputs uploaded) also stores the values")
    def _():
        fresh()
        load_example(case, browser=True)
        SS["page"] = PAGES[5]
        SS["mc_nreal"] = 10
        DOWNLOADS.clear()
        render(browser=True)
        lab = next(k for k in DOWNLOADS if k.startswith("1 · Download") and "input files" in k)
        UPLOADS["mc_out_up"] = [Upload("mc.zip", standin_outputs(zip_inputs(DOWNLOADS[lab][0])))]
        render(browser=True)
        res = next(iter(SS["mc_store"].values()))["results"]
        UPLOADS.pop("mc_out_up")
        assert res.loc[res["ok"] == True, "water_volume_above_invert_m3"].notna().all()  # noqa: E712
        SS["page"] = PAGES[6]
        SS["keep_cases_mc"] = True
        render(browser=True)
        c = _series_counts(_case_specs()[0])
        assert c.get("Monte Carlo", 0) >= 1, c
        return f"{c['Monte Carlo']} realisations plotted"


def main() -> int:
    global EXE
    if os.name == "nt":
        sys.exit("Linux and macOS only: the stand-in cannot run on Windows.")
    EXE = make_exe()
    for c in CASES:
        case_checks(c)
    project_checks()
    loading_checks()
    schematic_checks()
    overlay_checks("T1_water_overtopping")
    overlay_checks("T2_water_piping")
    security_checks()
    w = max(len(n) for _, n, _, _ in RESULTS)
    area = None
    for a, n, ok, d in RESULTS:
        if a != area:
            print(f"\n{a}")
            area = a
        tag = "SKIP" if ok is None else ("ok  " if ok else "FAIL")
        print(f"  {tag}  {n:<{w}}  {d}")
    n_fail = sum(1 for r in RESULTS if r[2] is False)
    n_skip = sum(1 for r in RESULTS if r[2] is None)
    print(f"\n{len(RESULTS) - n_fail - n_skip} of {len(RESULTS)} checks passed, "
          f"{n_fail} failed, {n_skip} skipped "
          "(stand-in executable; numbers are not DLBreach)")
    return 1 if n_fail else 0


if __name__ == "__main__":
    sys.exit(main())
